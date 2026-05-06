"""HTTP-level smoke test against the running backend."""
from __future__ import annotations
import json
import urllib.request
import urllib.parse
import os
import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment

BASE = "http://127.0.0.1:8765"


def _post_json(path: str, payload: dict) -> tuple[int, bytes, dict]:
    req = urllib.request.Request(
        BASE + path,
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    with urllib.request.urlopen(req) as resp:
        body = resp.read()
        return resp.status, body, dict(resp.headers)


def _post_multipart(path: str, file_path: Path) -> tuple[int, dict]:
    boundary = "----WebKitFormBoundary7MA4YWxkTrZu0gW"
    body = b""
    body += f"--{boundary}\r\n".encode()
    body += f'Content-Disposition: form-data; name="file"; filename="{file_path.name}"\r\n'.encode()
    body += b"Content-Type: application/octet-stream\r\n\r\n"
    body += file_path.read_bytes()
    body += f"\r\n--{boundary}--\r\n".encode()

    req = urllib.request.Request(
        BASE + path, data=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        method="POST",
    )
    with urllib.request.urlopen(req) as resp:
        return resp.status, json.loads(resp.read())


def build_sample(path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "תיקים"
    ws["A1"] = "משרד דוגמה"
    ws.merge_cells("A1:D1")
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = ["מס' תיק", "שם לקוח", "תאריך פתיחה", "סכום (₪)", "סטטוס", "הערות"]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=3, column=i, value=h)

    data = [
        (101, "כהן בע\"מ", "2024-03-01", 125000, "פתוח", "המתנה לדיון"),
        (102, "פרץ אלון",  "2024-04-15", 32000,  "סגור", "פסק דין"),
        (103, "ביטוח אחים","2024-05-20", 480000, "פתוח", "ערעור"),
        (104, "שלמה גולן", "2024-06-02", 78000,  "פתוח", "המתנה לחקירה"),
        (105, "נדל\"ן יהל","2024-06-30", 1200000,"פתוח", "סכסוך מקרקעין"),
    ]
    for ri, row in enumerate(data, start=4):
        for ci, val in enumerate(row, start=1):
            ws.cell(row=ri, column=ci, value=val)
    wb.save(path)
    wb.close()


def main() -> None:
    print("[health]", urllib.request.urlopen(BASE + "/api/health").read().decode())
    print("[templates]", json.loads(urllib.request.urlopen(BASE + "/api/templates").read())["templates"][0]["id"])

    tmp = Path(tempfile.mkdtemp(prefix="sxe_http_"))
    src = tmp / "data.xlsx"
    build_sample(str(src))

    status, info = _post_multipart("/api/upload", src)
    print(f"[upload] status={status}, file_id={info['file_id']}, ocr_used={info['ocr_used']}")
    file_id = info["file_id"]
    sheet = info["analysis"]["sheets"][0]
    print(f"  detected_header_row={sheet['detected_header_row']}")
    keep = [c["index"] for c in sheet["columns"] if not c["is_fully_empty"]]

    base_payload = {
        "file_id": file_id,
        "sheets": [{
            "sheet_name": sheet["sheet_name"],
            "column_indices": keep,
            "priorities": {str(c["index"]): c["suggested_priority"] for c in sheet["columns"] if c["index"] in keep},
            "filters": [
                {"column_index": next(c["index"] for c in sheet["columns"] if "סטטוס" in c["header"]),
                 "op": "eq", "value": "פתוח"}
            ],
            "skip_empty_rows": True,
            "header_row": None,
        }],
        "page_size": "A4",
        "orientation": "portrait",
        "margin_mm": 12,
        "rtl": True,
        "locale": "he",
        "smart_layout": True,
        "anchor_column_index": keep[0],
        "header": {
            "title": "דוח תיקים פעילים",
            "subtitle": "ייצוא אוטומטי",
            "logo_data_url": None,
            "show_page_numbers": True,
            "show_generated_date": True,
        },
        "explain": True,
    }

    # /preview/data
    pv_payload = {**base_payload, "row_limit": 50}
    status, body, _ = _post_json("/api/preview/data", pv_payload)
    js = json.loads(body)
    print(f"[preview/data] status={status}, sections={len(js['sections'])}, rows={js['sections'][0]['row_count']}")
    assert js["sections"][0]["row_count"] == 4  # only 4 'פתוח' rows after filtering

    # /preview/pdf
    status, body, _ = _post_json("/api/preview/pdf", pv_payload)
    print(f"[preview/pdf] status={status}, bytes={len(body)}, pdf={body[:4]==b'%PDF'}")

    # /export/pdf
    status, body, hdrs = _post_json("/api/export/pdf", base_payload)
    print(f"[export/pdf] status={status}, bytes={len(body)}, cd={hdrs.get('Content-Disposition')}")
    assert body[:4] == b"%PDF"

    # /export/word
    status, body, hdrs = _post_json("/api/export/word", base_payload)
    print(f"[export/word] status={status}, bytes={len(body)}, cd={hdrs.get('Content-Disposition')}")
    assert body[:2] == b"PK"

    # /auto/plan + /auto/preview/pdf + /auto/export/pdf — One-Click flow
    auto_payload = {
        "file_id": file_id,
        "locale": "he",
        "tier": "free",
        "watermark": {"preset": "none", "text": "", "opacity": 0.18, "color": "#888888"},
        "pdf_a": False,
        "row_limit": 50,
    }
    status, body, _ = _post_json("/api/auto/plan", auto_payload)
    js = json.loads(body)
    print(f"[auto/plan] status={status}, type={js['detection']['type']}, conf={js['detection']['confidence']:.2f}")
    assert js["detection"]["type"] in ("legal", "general", "financial", "inventory")
    assert js["plan"]["sheet_name"]

    status, body, _ = _post_json("/api/auto/preview/pdf", auto_payload)
    print(f"[auto/preview/pdf] status={status}, bytes={len(body)}, pdf={body[:4]==b'%PDF'}")
    assert body[:4] == b"%PDF"
    # Free tier should force a watermark — verify a watermark string is in the PDF
    assert b"Smart Excel Engine" in body, "free tier must imprint Smart Excel Engine watermark"

    # Pro tier — no forced watermark, allows PDF/A
    pro_payload = {
        **auto_payload,
        "tier": "pro",
        "watermark": {"preset": "none", "text": "", "opacity": 0.18, "color": "#888888"},
        "pdf_a": True,
    }
    status, body, _ = _post_json("/api/auto/preview/pdf", pro_payload)
    print(f"[auto/preview/pdf pro+pdfa] status={status}, bytes={len(body)}")
    assert body[:4] == b"%PDF"
    assert b"PDF/A" in body, "pro+pdf_a should advertise PDF/A best-effort"

    status, body, hdrs = _post_json("/api/auto/export/pdf", pro_payload)
    print(f"[auto/export/pdf pro] status={status}, bytes={len(body)}")
    print(f"     X-PDFA-Status={hdrs.get('X-PDFA-Status') or hdrs.get('x-pdfa-status')}")
    assert body[:4] == b"%PDF"
    pdfa_status = hdrs.get("X-PDFA-Status") or hdrs.get("x-pdfa-status")
    assert pdfa_status in ("verified", "best_effort", "failed"), f"unexpected pdfa status: {pdfa_status}"

    # Output Mode
    court_payload = {**pro_payload, "output_mode": "court_ready"}
    status, body, hdrs = _post_json("/api/auto/preview/pdf", court_payload)
    print(f"[auto/preview/pdf court_ready] status={status}, bytes={len(body)}")
    assert body[:4] == b"%PDF"

    draft_payload = {**auto_payload, "output_mode": "draft", "tier": "free"}
    status, body, _ = _post_json("/api/auto/preview/pdf", draft_payload)
    print(f"[auto/preview/pdf draft free] status={status}, bytes={len(body)}")
    assert body[:4] == b"%PDF"
    # free + draft → should imprint Smart Excel Engine still (free tier override)
    assert b"Smart Excel Engine" in body or b"DRAFT" in body, "expected free draft watermark"

    # Capabilities
    caps = json.loads(urllib.request.urlopen(BASE + "/api/capabilities").read())
    print(f"[capabilities] ocr={caps['ocr']}, pdfa_real={caps['pdfa_real']}")

    print("\n[OK] HTTP smoke test passed.")


if __name__ == "__main__":
    main()
