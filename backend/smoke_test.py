"""End-to-end smoke test for v2 architecture.

Exercises:
- Analyzer (header detection, fully/mostly empty, suggested priority)
- Merged-cell expansion (extract_table)
- Filters
- Layout engine (single-group + split into multiple page-groups)
- PDF renderer (multi-section with merge)
- Word renderer
- Templates loader
- OCR availability probe
"""
from __future__ import annotations
import tempfile
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from services.analyzer import analyze_workbook, extract_table
from services.filters import apply_filters
from services.layout_engine import build_columns, plan_layout
from services.pdf_renderer import render_pdf
from services.word_renderer import render_docx
from services.templates import list_templates
from services.detector import detect_report_type
from services import ocr, pdfa
from reportlab.lib.units import mm
from reportlab.lib.pagesizes import A4, portrait


def build_messy_workbook(path: str) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "תיקים"

    # Decorative rows
    ws["A1"] = "משרד עורכי דין דוגמה"
    ws["A2"] = "דוח תיקים פעילים"
    # Merge A1:E1 so we can verify merged-cell handling later (not in data range).
    ws.merge_cells("A1:E1")
    ws["A1"].alignment = Alignment(horizontal="center")

    headers = [
        "מס' תיק",       # high
        "שם לקוח",        # high
        "",                # empty header
        "סוג ההליך",      # medium
        "ערכאה",          # medium
        "תאריך פתיחה",    # high
        "סכום (₪)",       # high
        "",                # fully empty
        "סטטוס",          # high
        "הערות מורחבות",  # low (long text)
        "",                # mostly empty (1 value)
        "אחראי",           # medium
    ]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=3, column=i, value=h)

    base = [
        (101, "כהן בע\"מ",  None, "תביעה כספית",   "מחוזי תל אביב", "2024-03-01", 125000,  None, "פתוח", "המתנה לדיון בערכאה ראשונה — הצדדים הסכימו לדחיית מועד.", None, "עו\"ד לוי"),
        (102, "פרץ אלון",    None, "תביעה אזרחית",  "שלום ירושלים",  "2024-04-15", 32000,   None, "סגור", "ניתן פסק דין סופי.", None, "עו\"ד כהן"),
        (103, "ביטוח אחים",  None, "ערעור",         "עליון",         "2024-05-20", 480000,  None, "פתוח", "", None, "עו\"ד לוי"),
        (104, "שלמה גולן",   None, "תביעה כספית",   "מחוזי חיפה",    "2024-06-02", 78000,   None, "פתוח", "המתנה לחקירת חוקר פרטי.", None, "עו\"ד אבן"),
        (105, "נדל\"ן יהל",  None, "סכסוך מקרקעין", "מחוזי תל אביב", "2024-06-30", 1200000, None, "פתוח", "", None, "עו\"ד כהן"),
    ]
    row_idx = 4
    for cycle in range(5):
        for tup in base:
            for col_idx, val in enumerate(tup, start=1):
                ws.cell(row=row_idx, column=col_idx, value=val)
            row_idx += 1
        row_idx += 1  # one empty row each cycle

    # Merge a vertical group in the data: same case spread over 2 rows
    ws.merge_cells("A4:A5")  # 'מס' תיק' merged across 2 rows
    ws.cell(row=10, column=11, value="הערה חד פעמית")  # makes col K mostly_empty

    # Add a second sheet for sheet-merging test
    ws2 = wb.create_sheet("סיכום")
    ws2["A1"] = "סטטוס"
    ws2["B1"] = "כמות"
    ws2["A2"], ws2["B2"] = "פתוח", 4
    ws2["A3"], ws2["B3"] = "סגור", 1

    wb.save(path)
    wb.close()


def main() -> None:
    print(f"[OCR available]: {ocr.is_available()}")

    tmpdir = Path(tempfile.mkdtemp(prefix="sxe_"))
    src = tmpdir / "messy.xlsx"
    build_messy_workbook(str(src))
    print(f"[1] Sample built: {src}")

    analysis = analyze_workbook(str(src))
    sheets = analysis["sheets"]
    sh0 = sheets[0]
    print(f"[2] Sheet '{sh0['sheet_name']}': header_row={sh0['detected_header_row']}, "
          f"merged_ranges={len(sh0['merged_ranges'])}, columns={sh0['total_columns']}")
    assert sh0["detected_header_row"] == 3
    assert sh0["has_merged_cells"]

    fully_empty = [c for c in sh0["columns"] if c["is_fully_empty"]]
    print(f"    fully_empty cols: {len(fully_empty)} (expect 2)")
    assert len(fully_empty) >= 2

    # Show suggested priorities
    sugg = {c["header"]: c["suggested_priority"] for c in sh0["columns"]}
    print(f"    suggested priorities: {sugg}")

    keep = [c["index"] for c in sh0["columns"] if not c["is_fully_empty"]]
    table = extract_table(str(src), sh0["sheet_name"], keep, skip_empty_rows=True, expand_merged=True)
    print(f"[3] Extracted: {len(table['rows'])} rows × {len(table['headers'])} cols")
    # First two data rows should both have value '101' in column A (merged-cell expansion).
    assert table["rows"][0][0] == "101"
    assert table["rows"][1][0] == "101", "merged cell value should propagate to row 2"

    # --- Test filters: status='פתוח'
    status_pos = next(i for i, h in enumerate(table["headers"]) if "סטטוס" in h)
    status_col_idx = keep[status_pos]
    filtered = apply_filters(
        table["headers"], table["rows"], keep,
        [{"column_index": status_col_idx, "op": "eq", "value": "פתוח"}],
    )
    print(f"[4] After filter status=='פתוח': {len(filtered)} rows")
    assert all(r[status_pos] == "פתוח" for r in filtered)
    assert len(filtered) < len(table["rows"])

    # --- Test layout engine: build with widths + priorities, force narrow page to trigger split
    est_widths = {ci: 12 for ci in keep}  # placeholder
    for pos, ci in enumerate(keep):
        h = table["headers"][pos]
        for r in table["rows"][:50]:
            if pos < len(r):
                est_widths[ci] = max(est_widths[ci], min(60, len(r[pos]) + 2))
        est_widths[ci] = max(est_widths[ci], min(60, len(h) + 2))

    priorities = {}
    for c in sh0["columns"]:
        if c["index"] in keep:
            priorities[c["index"]] = c["suggested_priority"]

    cols = build_columns(table["headers"], keep, est_widths, priorities)

    # Wide page → single group
    wide = portrait(A4)[0] - 2 * 12 * mm
    layout_wide = plan_layout(columns=cols, available_width_pt=wide * 2, smart_layout=True, anchor_index=keep[0], locale="he")
    print(f"[5a] Wide layout: {len(layout_wide.groups)} group(s) (expect 1)")
    assert len(layout_wide.groups) == 1

    # Narrow page → split
    layout_split = plan_layout(columns=cols, available_width_pt=wide * 0.5, smart_layout=True, anchor_index=keep[0], locale="he")
    print(f"[5b] Narrow layout: {len(layout_split.groups)} groups (expect ≥ 2)")
    print(f"    explanations:")
    for line in layout_split.explanations[:5]:
        print(f"      - {line}")
    assert len(layout_split.groups) >= 2
    # Anchor column repeats in every group
    anchor_idx = keep[0]
    for g in layout_split.groups:
        assert any(c.index == anchor_idx for c in g.columns), "anchor missing from a group"

    # --- Render PDF (single sheet, merged file) using NORMAL layout
    section = {
        "title": sh0["sheet_name"],
        "layout": layout_wide,
        "rows": filtered,
        "positions_in_row": list(range(len(keep))),
    }
    pdf = render_pdf(
        sections=[section],
        page_size="A4", orientation="landscape", margin_mm=12,
        rtl=True, locale="he",
        header_title="דוח תיקים פעילים",
        header_subtitle="ייצוא אוטומטי — מצב פתוח בלבד",
        logo_data_url=None, show_page_numbers=True, show_generated_date=True,
    )
    pdf_path = tmpdir / "out.pdf"
    pdf_path.write_bytes(pdf)
    print(f"[6] PDF rendered: {len(pdf):,} bytes -> {pdf_path}")
    assert pdf[:4] == b"%PDF" and len(pdf) > 5000

    # --- Render PDF with SPLIT layout, then DOCX
    section_split = {
        "title": sh0["sheet_name"] + " (split)",
        "layout": layout_split,
        "rows": filtered,
        "positions_in_row": list(range(len(keep))),
    }
    pdf2 = render_pdf(
        sections=[section_split],
        page_size="A4", orientation="portrait", margin_mm=12,
        rtl=True, locale="he",
        header_title="דוח תיקים פעילים — פיצול",
        header_subtitle="", logo_data_url=None,
        show_page_numbers=True, show_generated_date=True,
    )
    (tmpdir / "out_split.pdf").write_bytes(pdf2)
    print(f"[6b] Split PDF: {len(pdf2):,} bytes")
    assert pdf2[:4] == b"%PDF"

    docx = render_docx(
        sections=[section],
        page_size="A4", orientation="landscape", margin_mm=12,
        rtl=True, locale="he",
        header_title="דוח תיקים פעילים", header_subtitle="",
        logo_data_url=None, show_page_numbers=True, show_generated_date=True,
    )
    (tmpdir / "out.docx").write_bytes(docx)
    print(f"[7] DOCX rendered: {len(docx):,} bytes")
    assert docx[:2] == b"PK"

    # Templates
    tpls = list_templates()
    print(f"[8] Templates loaded: {[t['id'] for t in tpls]}")
    assert any(t["id"] == "financial" for t in tpls)
    assert any(t["id"] == "general" for t in tpls)
    assert any(t["id"] == "legal" for t in tpls)
    assert any(t["id"] == "clean" for t in tpls)

    # Smart Report Detection — should classify the legal sheet
    detection = detect_report_type(analysis)
    print(f"[9] Detection: type={detection.detected_type} confidence={detection.confidence}")
    assert detection.detected_type == "legal", f"expected 'legal', got {detection.detected_type}"

    # Watermark + PDF/A flags
    pdf_wm = render_pdf(
        sections=[section],
        page_size="A4", orientation="landscape", margin_mm=12,
        rtl=True, locale="he",
        header_title="עם סימן מים", header_subtitle="",
        watermark_preset="draft", watermark_text="", watermark_opacity=0.18, watermark_color="#888888",
        pdf_a=True,
    )
    (tmpdir / "out_watermark_pdfa.pdf").write_bytes(pdf_wm)
    assert pdf_wm[:4] == b"%PDF"
    # PDF/A best-effort: check producer string in metadata
    assert b"PDF/A" in pdf_wm, "PDF should advertise PDF/A best-effort in producer metadata"
    assert b"Smart Excel Engine" in pdf_wm, "PDF should advertise the producer"
    print(f"[10] Watermark+PDF/A render: {len(pdf_wm):,} bytes (producer/PDF-A markers found)")

    # PDF/A real (via Ghostscript) — graceful degrade if not installed
    print(f"[11] Ghostscript available: {pdfa.is_available()}")
    out_bytes, status, msg_he, _ = pdfa.to_pdfa(pdf)
    print(f"     to_pdfa: status={status}, msg_he={msg_he[:80]}")
    if pdfa.is_available():
        assert status == "verified", f"with gs installed, expected 'verified', got {status}"
    else:
        assert status == "best_effort", f"without gs, expected 'best_effort', got {status}"

    # Output mode mapping
    from models.schemas import output_mode_defaults
    assert output_mode_defaults("draft") == {"watermark_preset": "draft", "pdf_a": False, "header_style": "standard"}
    assert output_mode_defaults("official") == {"watermark_preset": "none", "pdf_a": True, "header_style": "standard"}
    assert output_mode_defaults("court_ready") == {"watermark_preset": "none", "pdf_a": True, "header_style": "formal"}
    print("[12] Output mode mapping ok (draft/official/court_ready)")

    print("\n[OK] All tests passed.")
    print(f"Outputs: {tmpdir}")


if __name__ == "__main__":
    main()
