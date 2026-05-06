"""Excel analyzer with merged-cell support and priority hints."""
from __future__ import annotations
import openpyxl
from openpyxl.utils import get_column_letter
from dataclasses import dataclass, asdict
from typing import Any


@dataclass
class ColumnInfo:
    index: int
    letter: str
    header: str
    detected_type: str
    total_cells: int
    non_empty_cells: int
    fill_ratio: float
    is_fully_empty: bool
    is_mostly_empty: bool
    sample_values: list[Any]
    estimated_width: int
    suggested_priority: str  # high|medium|low


@dataclass
class MergedRange:
    min_row: int
    max_row: int
    min_col: int
    max_col: int


@dataclass
class SheetAnalysis:
    sheet_name: str
    detected_header_row: int
    data_start_row: int
    data_end_row: int
    total_rows: int
    total_columns: int
    columns: list[ColumnInfo]
    has_merged_cells: bool
    merged_ranges: list[MergedRange]
    data_density: float


# Heuristic: column-name keyword stems suggesting "important" identifiers.
HIGH_PRIORITY_HINTS = {
    "id", "code", "מספר", "מס'", "מס.", "מס'",
    "שם", "name", "כותרת", "title",
    "תאריך", "date",
    "סכום", "amount", "total", "סה\"כ", "סהכ",
    "סטטוס", "status",
}
LOW_PRIORITY_HINTS = {
    "הערה", "הערות", "note", "notes", "comment", "comments", "remark", "remarks",
    "תיאור", "description", "details", "פרטים",
}


def _classify(values: list[Any]) -> str:
    if not values:
        return "empty"
    types = set()
    for v in values:
        if v is None:
            continue
        if isinstance(v, bool):
            types.add("bool")
        elif isinstance(v, (int, float)):
            types.add("number")
        else:
            try:
                float(str(v).replace(",", "").replace("%", ""))
                types.add("number")
            except (ValueError, AttributeError):
                types.add("text")
    if not types:
        return "empty"
    if types == {"number"}:
        return "number"
    if types == {"bool"}:
        return "bool"
    return "text"


def _detect_header_row(ws, max_scan: int = 15) -> int:
    best_row = 1
    best_score = -1.0
    rows_to_scan = min(max_scan, ws.max_row or 1)
    for row_idx in range(1, rows_to_scan + 1):
        rows = list(ws.iter_rows(min_row=row_idx, max_row=row_idx, values_only=True))
        if not rows:
            continue
        cells = rows[0]
        non_empty = [c for c in cells if c is not None and str(c).strip() != ""]
        if not non_empty:
            continue
        text_count = sum(1 for c in non_empty if not isinstance(c, (int, float)) or isinstance(c, bool))
        fill_ratio = len(non_empty) / max(1, len(cells))
        text_ratio = text_count / max(1, len(non_empty))
        score = fill_ratio * 0.6 + text_ratio * 0.4 - row_idx * 0.001
        if score > best_score:
            best_score = score
            best_row = row_idx
    return best_row


def _estimate_width(values: list[Any], header: str) -> int:
    longest = max([len(str(v)) for v in values if v is not None] + [len(header or "")], default=8)
    return min(60, max(8, longest + 2))


def _suggest_priority(header: str, fill_ratio: float, dtype: str, col_index: int, est_width: int) -> str:
    h = (header or "").lower().strip()

    # Rule 1: known high-priority keywords
    for kw in HIGH_PRIORITY_HINTS:
        if kw.lower() in h:
            return "high"

    # Rule 2: known low-priority keywords
    for kw in LOW_PRIORITY_HINTS:
        if kw.lower() in h:
            return "low"

    # Rule 3: leftmost columns are usually identifiers
    if col_index == 0:
        return "high"

    # Rule 4: very wide free-text columns are usually low-priority detail
    if est_width >= 40 and dtype == "text":
        return "low"

    # Rule 5: well-filled columns are more important than sparse ones
    if fill_ratio < 0.2:
        return "low"
    if fill_ratio >= 0.85:
        return "medium"

    return "medium"


def analyze_workbook(path: str) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = []
    for ws in wb.worksheets:
        sheets.append(asdict(_analyze_sheet(ws)))
    wb.close()
    return {"sheets": sheets}


def _analyze_sheet(ws) -> SheetAnalysis:
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    if max_row == 0 or max_col == 0:
        return SheetAnalysis(
            sheet_name=ws.title,
            detected_header_row=1,
            data_start_row=2,
            data_end_row=1,
            total_rows=0,
            total_columns=0,
            columns=[],
            has_merged_cells=False,
            merged_ranges=[],
            data_density=0.0,
        )

    header_row = _detect_header_row(ws)
    data_start = header_row + 1
    data_end = max_row

    all_rows: list[tuple] = list(
        ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True)
    )
    header_cells = all_rows[header_row - 1] if len(all_rows) >= header_row else tuple()

    total_data_cells = 0
    non_empty_data_cells = 0
    columns: list[ColumnInfo] = []

    for col_idx in range(max_col):
        header_val = header_cells[col_idx] if col_idx < len(header_cells) else None
        header_text = str(header_val).strip() if header_val is not None else ""

        column_values: list[Any] = []
        for row_idx in range(data_start - 1, data_end):
            if row_idx >= len(all_rows):
                break
            row = all_rows[row_idx]
            v = row[col_idx] if col_idx < len(row) else None
            column_values.append(v)

        non_empty = [v for v in column_values if v is not None and str(v).strip() != ""]
        total = len(column_values)
        fill = (len(non_empty) / total) if total else 0.0
        total_data_cells += total
        non_empty_data_cells += len(non_empty)

        body_empty = len(non_empty) == 0
        header_empty = header_text == ""
        dtype = _classify(non_empty)
        est_width = _estimate_width(non_empty, header_text)
        suggested = _suggest_priority(header_text, fill, dtype, col_idx, est_width)

        columns.append(
            ColumnInfo(
                index=col_idx,
                letter=get_column_letter(col_idx + 1),
                header=header_text or f"עמודה {col_idx + 1}",
                detected_type=dtype,
                total_cells=total,
                non_empty_cells=len(non_empty),
                fill_ratio=round(fill, 3),
                is_fully_empty=body_empty and header_empty,
                is_mostly_empty=fill < 0.1 and not body_empty,
                sample_values=[str(v) if v is not None else "" for v in non_empty[:5]],
                estimated_width=est_width,
                suggested_priority=suggested,
            )
        )

    density = (non_empty_data_cells / total_data_cells) if total_data_cells else 0.0
    merged = [
        MergedRange(
            min_row=r.min_row, max_row=r.max_row,
            min_col=r.min_col, max_col=r.max_col,
        )
        for r in ws.merged_cells.ranges
    ]

    return SheetAnalysis(
        sheet_name=ws.title,
        detected_header_row=header_row,
        data_start_row=data_start,
        data_end_row=data_end,
        total_rows=max_row,
        total_columns=max_col,
        columns=columns,
        has_merged_cells=len(merged) > 0,
        merged_ranges=merged,
        data_density=round(density, 3),
    )


def _expand_merged(ws, max_row: int, max_col: int) -> list[list[Any]]:
    """Read all values, but for merged ranges, copy the top-left value to every cell in the range."""
    grid: list[list[Any]] = [[None] * max_col for _ in range(max_row)]
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True)):
        for c_idx, v in enumerate(row):
            grid[r_idx][c_idx] = v
    for mr in ws.merged_cells.ranges:
        r0, c0 = mr.min_row - 1, mr.min_col - 1
        if r0 >= max_row or c0 >= max_col:
            continue
        anchor = grid[r0][c0]
        for r in range(mr.min_row - 1, min(mr.max_row, max_row)):
            for c in range(mr.min_col - 1, min(mr.max_col, max_col)):
                if grid[r][c] is None:
                    grid[r][c] = anchor
    return grid


def extract_table(
    path: str,
    sheet_name: str,
    column_indices: list[int],
    skip_empty_rows: bool = True,
    header_row: int | None = None,
    expand_merged: bool = True,
) -> dict:
    """Returns {'headers': [...], 'rows': [[...], ...]}"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet '{sheet_name}' not found")
    ws = wb[sheet_name]

    if header_row is None:
        header_row = _detect_header_row(ws)

    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    if expand_merged and ws.merged_cells.ranges:
        grid = _expand_merged(ws, max_row, max_col)
    else:
        grid = [
            list(row) for row in ws.iter_rows(
                min_row=1, max_row=max_row, max_col=max_col, values_only=True,
            )
        ]
    wb.close()

    if not grid:
        return {"headers": [], "rows": []}

    header_cells = grid[header_row - 1] if header_row - 1 < len(grid) else []
    headers: list[str] = []
    for ci in column_indices:
        if ci < len(header_cells) and header_cells[ci] is not None:
            headers.append(str(header_cells[ci]).strip() or f"עמודה {ci + 1}")
        else:
            headers.append(f"עמודה {ci + 1}")

    out_rows: list[list[str]] = []
    for r_idx in range(header_row, len(grid)):
        row = grid[r_idx]
        out: list[str] = []
        for ci in column_indices:
            v = row[ci] if ci < len(row) else None
            if v is None:
                out.append("")
            elif isinstance(v, float) and v.is_integer():
                out.append(str(int(v)))
            else:
                out.append(str(v))
        if skip_empty_rows and not any(s.strip() for s in out):
            continue
        out_rows.append(out)

    return {"headers": headers, "rows": out_rows}
