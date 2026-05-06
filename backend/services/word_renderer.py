"""Word renderer aligned with the layout engine."""
from __future__ import annotations
from io import BytesIO
import base64
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .layout_engine import LayoutResult, slice_rows


def _set_cell_shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _set_cell_borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "BFBFBF")
        borders.append(b)
    tc_pr.append(borders)


def _set_rtl_paragraph(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    p_pr.append(bidi)


def _set_table_rtl(table) -> None:
    bidi = OxmlElement("w:bidiVisual")
    table._tbl.tblPr.append(bidi)


def _add_page_x_of_y_footer(doc, locale: str) -> None:
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = "עמוד " if locale.startswith("he") else "Page "
    p.add_run(label)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    p.add_run()._r.append(fld_begin)
    instr = OxmlElement("w:instrText")
    instr.text = "PAGE"
    p.add_run()._r.append(instr)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    p.add_run()._r.append(fld_end)
    p.add_run(" / " if locale.startswith("he") else " of ")
    fld_begin2 = OxmlElement("w:fldChar")
    fld_begin2.set(qn("w:fldCharType"), "begin")
    p.add_run()._r.append(fld_begin2)
    instr2 = OxmlElement("w:instrText")
    instr2.text = "NUMPAGES"
    p.add_run()._r.append(instr2)
    fld_end2 = OxmlElement("w:fldChar")
    fld_end2.set(qn("w:fldCharType"), "end")
    p.add_run()._r.append(fld_end2)


def render_docx(
    *,
    sections: list[dict],
    page_size: str = "A4",
    orientation: str = "portrait",
    margin_mm: float = 12,
    rtl: bool = True,
    locale: str = "he",
    header_title: str = "",
    header_subtitle: str = "",
    logo_data_url: str | None = None,
    show_page_numbers: bool = True,
    show_generated_date: bool = True,
) -> bytes:
    doc = Document()
    section = doc.sections[0]

    if page_size.upper() == "A3":
        width, height = Cm(29.7), Cm(42.0)
    else:
        width, height = Cm(21.0), Cm(29.7)
    if orientation.lower() == "landscape":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = height, width
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = width, height
    margin = Cm(margin_mm / 10)
    section.top_margin = section.bottom_margin = margin
    section.left_margin = section.right_margin = margin

    # Document header
    if logo_data_url and logo_data_url.startswith("data:"):
        try:
            _, b64 = logo_data_url.split(",", 1)
            from io import BytesIO as _B
            img = _B(base64.b64decode(b64))
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(img, width=Inches(1.2))
        except Exception:
            pass

    if header_title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if rtl:
            _set_rtl_paragraph(p)
        run = p.add_run(header_title)
        run.bold = True
        run.font.size = Pt(16)
        run.font.name = "Arial"

    if header_subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if rtl:
            _set_rtl_paragraph(p)
        run = p.add_run(header_subtitle)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        run.font.name = "Arial"

    if show_generated_date:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if rtl:
            _set_rtl_paragraph(p)
        stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
        label = f"נוצר: {stamp}" if locale.startswith("he") else f"Generated: {stamp}"
        run = p.add_run(label)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
        run.font.name = "Arial"

    if show_page_numbers:
        _add_page_x_of_y_footer(doc, locale)

    # Sections + groups
    for s_idx, sec in enumerate(sections):
        if s_idx > 0:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

        title = sec.get("title")
        layout: LayoutResult = sec["layout"]
        rows: list[list[str]] = sec["rows"]
        positions_in_row: list[int] = sec["positions_in_row"]

        if title:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if rtl:
                _set_rtl_paragraph(p)
            run = p.add_run(title)
            run.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

        for gi, group in enumerate(layout.groups):
            if gi > 0:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

            if len(layout.groups) > 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if rtl:
                    _set_rtl_paragraph(p)
                txt = (
                    f"קבוצת עמודים {gi + 1} מתוך {len(layout.groups)}"
                    if locale.startswith("he")
                    else f"Page-group {gi + 1} of {len(layout.groups)}"
                )
                run = p.add_run(txt)
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

            cols = list(group.columns)
            sliced = slice_rows(rows, [positions_in_row[c.pos] for c in group.columns])

            if rtl:
                cols = list(reversed(cols))
                sliced = [list(reversed(r)) for r in sliced]

            table = doc.add_table(rows=1, cols=len(cols))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            if rtl:
                _set_table_rtl(table)

            hdr = table.rows[0].cells
            for i, c in enumerate(cols):
                cell = hdr[i]
                cell.text = ""
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if rtl:
                    _set_rtl_paragraph(p)
                run = p.add_run(c.header)
                run.bold = True
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.name = "Arial"
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                _set_cell_shade(cell, "1F4E79")
                _set_cell_borders(cell)

            for ri, row_data in enumerate(sliced):
                cells = table.add_row().cells
                shade = "FFFFFF" if ri % 2 == 0 else "F2F6FB"
                padded = list(row_data) + [""] * (len(cols) - len(row_data))
                for ci, val in enumerate(padded[: len(cols)]):
                    cell = cells[ci]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
                    if rtl:
                        _set_rtl_paragraph(p)
                    run = p.add_run(str(val) if val is not None else "")
                    run.font.size = Pt(10)
                    run.font.name = "Arial"
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    _set_cell_shade(cell, shade)
                    _set_cell_borders(cell)

            # Repeat header on every page
            tr = table.rows[0]._tr
            tr_pr = tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            tr_pr.append(th)

        if layout.explanations:
            p = doc.add_paragraph()
            if rtl:
                _set_rtl_paragraph(p)
            run = p.add_run("הסבר פריסה:" if locale.startswith("he") else "Layout notes:")
            run.bold = True
            run.font.size = Pt(10)
            for ex in layout.explanations:
                p = doc.add_paragraph()
                if rtl:
                    _set_rtl_paragraph(p)
                run = p.add_run("• " + ex)
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
